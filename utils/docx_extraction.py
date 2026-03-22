"""
DOCX Text Extraction — Vercel-compatible
==========================================
Extracts text from .docx: paragraphs, tables, embedded images via Google Vision API.
"""

import io, os
from docx import Document
from PIL import Image


def _extract_images(part):
    images = []
    try:
        for rel in part.rels.values():
            if "image" in rel.reltype:
                try:
                    images.append(Image.open(io.BytesIO(rel.target_part.blob)))
                except Exception:
                    pass
    except Exception:
        pass
    return images


def extract_text_from_docx(docx_path: str, ocr_lang: str = "en",
                            ocr_images: bool = True) -> str:
    if not docx_path.lower().endswith(".docx"):
        raise ValueError(f"Expected .docx, got: '{os.path.basename(docx_path)}'")
    try:
        doc = Document(docx_path)
    except Exception as e:
        raise RuntimeError(f"Could not open DOCX: {e}")

    parts = []

    for para in doc.paragraphs:
        line = para.text.strip()
        if line:
            parts.append(line)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    if ocr_images:
        from utils.ocr import run_ocr
        for img in _extract_images(doc.part):
            try:
                ocr_text = run_ocr(img, lang=ocr_lang, preprocess=True)
                if ocr_text:
                    parts.append(f"[IMAGE TEXT]\n{ocr_text}")
            except Exception:
                pass

    return "\n".join(parts).strip()
